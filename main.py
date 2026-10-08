from __future__ import annotations

import csv
import io
import json
import math
import re
import signal
import sys
import threading
from dataclasses import dataclass
from functools import lru_cache
from itertools import permutations
from pathlib import Path
from typing import Optional

try:
    import customtkinter as ctk
except ImportError:
    print("Установите зависимости: pip install customtkinter openpyxl Pillow cairosvg")
    raise

import tkinter as tk
from tkinter import filedialog, messagebox

try:
    from PIL import Image, ImageDraw, ImageFont, ImageTk
except ImportError:
    Image = None
    ImageDraw = None
    ImageFont = None
    ImageTk = None


ROOT = Path(__file__).resolve().parent
DB_PATH = ROOT / "cells_db.txt"
STATE_PATH = ROOT / "main_state.json"
LOGO_PATHS = [ROOT / "logo_RD.png", ROOT / "logo.png", ROOT / "logo_RD.svg", ROOT / "logo.svg"]
ICO_PATHS = [ROOT / "logo_RD_ico.ico", ROOT / "logo_RD.ico", ROOT / "logo.ico"]
ICON_PATHS = [*ICO_PATHS, ROOT / "logo_RD_ico.png", ROOT / "logo_RD.png"]
HEADER = [
    "наименование", "бренд", "тип_химии", "тип_формы", "длина_мм", "ширина_мм",
    "высота_мм", "емкость_Ач", "напряжение_ном_В", "напряжение_мин_В",
    "напряжение_макс_В", "макс_ток_разряда_А", "макс_ток_заряда_А",
    "мин_темп_разряда", "макс_темп_разряда", "масса_г",
    "внутр_сопротивление_мОм", "циклов", "емкость_после_1000_циклов_%",
]
BG = "#fdf6f4"
FIELD_BG = "#f6dcd6"
BUTTON = "#c9c9c9"
BUTTON_HOVER = "#b5b5b5"
ERROR = "#b00020"
OK = "#1b7a1b"
FONT = ("GOST type A", 14)


@dataclass
class Cell:
    name: str
    brand: str
    chemistry: str
    shape: str
    length: float
    width: float
    height: float
    capacity: float
    voltage_nom: float
    voltage_min: Optional[float]
    voltage_max: Optional[float]
    max_discharge: float
    max_charge: Optional[float]
    temp_min: Optional[float]
    temp_max: Optional[float]
    mass: Optional[float]
    resistance: Optional[float]
    cycles: Optional[float]
    capacity_1000: Optional[float]


@dataclass
class Pack:
    s: int
    p: int
    n: int
    dims: tuple[float, float, float]
    grid: tuple[int, int, int]
    orientation: tuple[str, str, str]
    energy: float
    capacity: float
    voltage: float
    max_current: float
    runtime: float
    volume: float
    ok_time: bool
    ok_dims: bool
    warning: str
    marks: str = ""


def ensure_db() -> None:
    if not DB_PATH.exists() or DB_PATH.stat().st_size == 0:
        DB_PATH.write_text(";".join(HEADER) + "\n", encoding="utf-8")


def to_float(value: str) -> Optional[float]:
    value = value.strip().replace(",", ".")
    if not value:
        return None
    return float(value)


def must_float(row: dict[str, str], key: str) -> float:
    value = to_float(row.get(key, ""))
    if value is None:
        raise ValueError(key)
    return value


def load_db() -> tuple[list[Cell], int]:
    ensure_db()
    cells: list[Cell] = []
    skipped = 0
    with DB_PATH.open("r", encoding="utf-8", newline="") as fh:
        reader = csv.reader(fh, delimiter=";")
        try:
            header = next(reader)
        except StopIteration:
            return [], 0
        if header != HEADER:
            print("Предупреждение: заголовок базы отличается от ожидаемого.")
        for line_no, parts in enumerate(reader, start=2):
            if not parts or all(not p.strip() for p in parts):
                continue
            if len(parts) != len(HEADER):
                skipped += 1
                print(f"Пропуск строки {line_no}: ожидалось {len(HEADER)} полей, получено {len(parts)}")
                continue
            row = {key: value.strip() for key, value in zip(HEADER, parts)}
            try:
                shape = row["тип_формы"].upper()
                if shape not in {"C", "P"}:
                    raise ValueError("тип_формы")
                cells.append(Cell(
                    row["наименование"], row["бренд"], row["тип_химии"], shape,
                    must_float(row, "длина_мм"), must_float(row, "ширина_мм"), must_float(row, "высота_мм"),
                    must_float(row, "емкость_Ач"), must_float(row, "напряжение_ном_В"),
                    to_float(row["напряжение_мин_В"]), to_float(row["напряжение_макс_В"]),
                    must_float(row, "макс_ток_разряда_А"), to_float(row["макс_ток_заряда_А"]),
                    to_float(row["мин_темп_разряда"]), to_float(row["макс_темп_разряда"]),
                    to_float(row["масса_г"]), to_float(row["внутр_сопротивление_мОм"]),
                    to_float(row["циклов"]), to_float(row["емкость_после_1000_циклов_%"]),
                ))
            except Exception as exc:
                skipped += 1
                print(f"Пропуск строки {line_no}: {exc}")
    print(f"База: {DB_PATH}")
    print(f"Загружено ячеек: {len(cells)}")
    print(f"Пропущено строк: {skipped}")
    return cells, skipped


def generated_rd_logo(size: int):
    if Image is None or ImageDraw is None:
        return None
    image = Image.new("RGBA", (size, size), (253, 246, 244, 255))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((1, 1, size - 2, size - 2), radius=max(4, size // 5), fill="#f6dcd6", outline="#1a1a1a", width=2)
    try:
        font = ImageFont.truetype("arialbd.ttf", max(14, int(size * 0.42))) if ImageFont else None
    except Exception:
        font = None
    text = "RD"
    box = draw.textbbox((0, 0), text, font=font)
    x = (size - (box[2] - box[0])) / 2
    y = (size - (box[3] - box[1])) / 2 - 1
    draw.text((x, y), text, fill="#1a1a1a", font=font)
    return ctk.CTkImage(light_image=image, dark_image=image, size=(size, size))


def add_black_background(image: Image.Image) -> Image.Image:
    image = image.convert("RGBA")
    bg = Image.new("RGBA", image.size, (0, 0, 0, 255))
    bg.alpha_composite(image)
    return bg


def svg_dimensions(path: Path) -> tuple[int, int]:
    text = path.read_text(encoding="utf-8", errors="ignore")[:1000]
    width = re.search(r'width="(\d+(?:\.\d+)?)px"', text)
    height = re.search(r'height="(\d+(?:\.\d+)?)px"', text)
    if width and height:
        return int(float(width.group(1))), int(float(height.group(1)))
    viewbox = re.search(r'viewBox="[^"]*?(\d+(?:\.\d+)?)\s+(\d+(?:\.\d+)?)"', text)
    if viewbox:
        return int(float(viewbox.group(1))), int(float(viewbox.group(2)))
    return 1, 1


def load_logo(size: int):
    if Image is None:
        print("Pillow не установлен, логотип отключен.")
        return None
    for path in LOGO_PATHS:
        if not path.exists():
            continue
        try:
            if path.suffix.lower() == ".svg":
                try:
                    import cairosvg
                    svg_w, svg_h = svg_dimensions(path)
                    out_w = max(size, min(180, int(size * svg_w / max(svg_h, 1))))
                    png = cairosvg.svg2png(url=str(path), output_width=out_w, output_height=size)
                    image = Image.open(io.BytesIO(png)).convert("RGBA")
                except Exception as exc:
                    print(f"Не удалось прочитать SVG {path.name}: {exc}")
                    continue
            else:
                source = Image.open(path).convert("RGBA")
                out_w = max(size, min(180, int(size * source.width / max(source.height, 1))))
                image = source.resize((out_w, size))
            image = add_black_background(image)
            return ctk.CTkImage(light_image=image, dark_image=image, size=image.size)
        except Exception as exc:
            print(f"Логотип {path.name} не загружен: {exc}")
    print("Логотип не найден или не прочитан, используется встроенный знак RD.")
    return generated_rd_logo(size)


def load_window_icon(master: tk.Misc, size: int) -> Optional[ImageTk.PhotoImage]:
    if Image is None or ImageTk is None:
        return None
    for path in ICON_PATHS:
        if not path.exists():
            continue
        try:
            source = Image.open(path).convert("RGBA")
            scale = min(size / max(source.width, 1), size / max(source.height, 1))
            logo = source.resize((max(1, int(source.width * scale)), max(1, int(source.height * scale))))
            image = Image.new("RGBA", (size, size), (0, 0, 0, 255))
            image.alpha_composite(logo, ((size - logo.width) // 2, (size - logo.height) // 2))
            return ImageTk.PhotoImage(image, master=master)
        except Exception as exc:
            print(f"Иконка {path.name} не загружена: {exc}")
    fallback = generated_rd_logo(size)
    if fallback is None:
        return None
    image = fallback.cget("light_image")
    return ImageTk.PhotoImage(image, master=master)


def fmt(value: Optional[float], unit: str = "", digits: int = 2) -> str:
    if value is None:
        return "—"
    text = f"{value:.{digits}f}".rstrip("0").rstrip(".")
    return f"{text} {unit}".strip()


def shape_name(shape: str) -> str:
    return "цилиндрическая" if shape.upper() == "C" else "призматическая"


def rounded_dims(dims: tuple[float, float, float]) -> tuple[float, float, float]:
    return tuple(round(x, 1) for x in dims)  # type: ignore[return-value]


def row_size(count: int, size: float, gap: float) -> float:
    return count * size + max(0, count - 1) * gap


MAX_RESULT_PACKS = 160
MAX_LAYOUTS_PER_SCHEME = 5
MAX_DIM_LAYOUTS_PER_SCHEME = 3


def effective_target_dims(target: tuple[float, float, float], cut: float, wall: float) -> tuple[float, float, float]:
    return tuple(max(1.0, d * (1 + cut / 100.0) - 2 * wall) for d in target)  # type: ignore[return-value]


def fill_stats(dims: tuple[float, float, float], effective: tuple[float, float, float]) -> tuple[float, float, float]:
    target_volume = max(1.0, effective[0] * effective[1] * effective[2])
    filled_volume = math.prod(min(dims[i], effective[i]) for i in range(3))
    empty_ratio = sum(max(0.0, effective[i] - dims[i]) / effective[i] for i in range(3))
    overflow_ratio = sum(max(0.0, dims[i] - effective[i]) / effective[i] for i in range(3))
    return filled_volume / target_volume, empty_ratio, overflow_ratio


def pack_fill_key(pack: Pack, effective: tuple[float, float, float]) -> tuple[object, ...]:
    fill_ratio, empty_ratio, overflow_ratio = fill_stats(pack.dims, effective)
    return (not pack.ok_time, -fill_ratio, empty_ratio, overflow_ratio, -pack.n, -pack.runtime, pack.s, pack.p, pack.grid)


@lru_cache(maxsize=None)
def grid_candidates(n: int) -> list[tuple[int, int, int]]:
    triples = []
    for c in range(1, n + 1):
        max_a = math.ceil(n / c)
        for a in range(1, max_a + 1):
            b = math.ceil(n / (a * c))
            if a * b * c >= n:
                triples.append((a, b, c))
    empty_limit = max(12, math.ceil(n * 0.12))
    filtered = [t for t in triples if t[0] * t[1] * t[2] - n <= empty_limit]
    return sorted(set(filtered), key=lambda t: (t[0] * t[1] * t[2] - n, max(t), sum(t)))[:180]


def layout_candidates(
    cell: Cell,
    n: int,
    consider_dims: bool,
    target: tuple[float, float, float],
    gap: float,
    wall: float,
    bms: float,
    bms_side: str,
    cut: float,
    limit: int = 1,
) -> list[tuple[tuple[float, float, float], tuple[int, int, int], tuple[str, str, str], bool, str]]:
    base_axes = {"L": cell.length, "W": cell.width, "H": cell.height}
    effective = effective_target_dims(target, cut, wall)
    all_items = []
    variants = []
    near = None
    for grid in grid_candidates(n):
        for orient in set(permutations(("L", "W", "H"))):
            sizes = [base_axes[axis] for axis in orient]
            dims = (
                row_size(grid[0], sizes[0], gap),
                row_size(grid[1], sizes[1], gap),
                row_size(grid[2], sizes[2], gap),
            )
            axis = {"front": 0, "side": 1, "top": 2}.get(bms_side, 0)
            dims_with_bms = list(dims)
            dims_with_bms[axis] += bms
            rdims = rounded_dims(tuple(dims_with_bms))
            overflow = [max(0.0, rdims[i] - effective[i]) for i in range(3)]
            volume = rdims[0] * rdims[1] * rdims[2]
            ok = not consider_dims or all(v <= 1e-9 for v in overflow)
            item = (volume, rdims, grid, orient, ok, overflow)
            all_items.append(item)
            if ok:
                variants.append(item)
            if near is None or sum(overflow) < sum(near[-1]):
                near = item
    sources = all_items if consider_dims else (variants if variants else ([near] if near is not None else []))
    if not sources:
        return []
    if consider_dims:
        sources = sorted(sources, key=lambda item: (-fill_stats(item[1], effective)[0], fill_stats(item[1], effective)[1], fill_stats(item[1], effective)[2], item[2], item[3]))
    else:
        sources = sorted(sources, key=lambda item: (item[0], max(item[2]), sum(item[2]), item[2], item[3]))
    picked = []
    if not consider_dims and len(sources) > 1:
        selectors = (
            min(sources, key=lambda item: item[0]),
            max(sources, key=lambda item: max(item[2])),
            min(sources, key=lambda item: max(item[2])),
            min(sources, key=lambda item: max(item[2]) - min(item[2])),
            max(sources, key=lambda item: sum(1 for count in item[2] if count > 1)),
        )
        for item in selectors:
            if item not in picked:
                picked.append(item)
    for item in sources:
        if item not in picked:
            picked.append(item)
        if len(picked) >= limit:
            break
    result = []
    for _volume, dims, grid, orient, ok, overflow in picked[:limit]:
        warning = ""
        if consider_dims and not ok:
            axes = ["глубине", "ширине", "высоте"]
            idx = max(range(3), key=lambda i: overflow[i])
            warning = f"⚠ превышение по {axes[idx]} на {overflow[idx]:.1f} мм"
        result.append((dims, grid, orient, ok, warning))
    return result


def best_layout(
    cell: Cell,
    n: int,
    consider_dims: bool,
    target: tuple[float, float, float],
    gap: float,
    wall: float,
    bms: float,
    bms_side: str,
    cut: float,
) -> tuple[tuple[float, float, float], tuple[int, int, int], tuple[str, str, str], bool, str]:
    return layout_candidates(cell, n, consider_dims, target, gap, wall, bms, bms_side, cut, 1)[0]


def calculate(
    cell: Cell,
    load_w: float,
    required_h: float,
    consider_dims: bool,
    target: tuple[float, float, float],
    cut: float,
    wall: float,
    gap: float,
    bms: float,
    bms_side: str,
) -> list[Pack]:
    packs: list[Pack] = []
    rejected_time: list[Pack] = []
    rejected_dims: list[Pack] = []
    e_cell = cell.voltage_nom * cell.capacity
    for s in range(1, 101):
        for p in range(1, 101):
            n = s * p
            if n > 400:
                continue
            energy = e_cell * n
            runtime = energy / load_w
            ok_time = runtime >= required_h
            layout_limit = MAX_DIM_LAYOUTS_PER_SCHEME if consider_dims else MAX_LAYOUTS_PER_SCHEME
            for dims, grid, orient, ok_dims, warning in layout_candidates(cell, n, consider_dims, target, gap, wall, bms, bms_side, cut, layout_limit):
                pack = Pack(
                    s, p, n, dims, grid, orient, energy, cell.capacity * p, cell.voltage_nom * s,
                    cell.max_discharge * p, runtime, dims[0] * dims[1] * dims[2], ok_time, ok_dims,
                    "" if ok_time and ok_dims else (warning or f"⚠ время меньше требуемого на {required_h - runtime:.2f} ч"),
                )
                if ok_time and ok_dims:
                    packs.append(pack)
                elif ok_time:
                    rejected_dims.append(pack)
                else:
                    rejected_time.append(pack)
    if consider_dims:
        effective = effective_target_dims(target, cut, wall)
        packs.extend(rejected_dims)
        packs.sort(key=lambda x: pack_fill_key(x, effective))
        rejected_dims = []
    else:
        packs.sort(key=lambda x: (x.n, x.volume, x.s))
    final = packs[:MAX_RESULT_PACKS]
    if len(final) < 3:
        rejected_time.sort(key=lambda x: abs(required_h - x.runtime))
        if consider_dims:
            rejected_dims.sort(key=lambda x: pack_fill_key(x, effective_target_dims(target, cut, wall)))
        else:
            rejected_dims.sort(key=lambda x: x.volume)
        final.extend(rejected_time[: max(0, 3 - len(final))])
        final.extend(rejected_dims[: max(0, 3 - len(final))])
    if not final and rejected_time:
        final = sorted(rejected_time, key=lambda x: abs(required_h - x.runtime))[:1]
    if consider_dims:
        final = sorted(final[:MAX_RESULT_PACKS], key=lambda x: pack_fill_key(x, effective_target_dims(target, cut, wall)))
    else:
        final = sorted(final[:MAX_RESULT_PACKS], key=lambda x: (x.n, x.s, x.p, x.grid, x.volume))
    if final:
        max_s = max(final, key=lambda x: x.s)
        max_p = max(final, key=lambda x: x.p)
        balanced = min(final, key=lambda x: abs(x.s - x.p))
        for pack, mark in ((max_s, "S"), (max_p, "P"), (balanced, "B")):
            pack.marks = "".join(sorted(set(pack.marks + mark)))
    return final


class MainApp(ctk.CTk):
    def __init__(self) -> None:
        super().__init__()
        ctk.set_appearance_mode("light")
        self.title("RD конфигуратор АКБ")
        self.configure(fg_color=BG)
        self.minsize(1000, 700)
        self._center(0.7, 0.7)
        self.logo = load_logo(40)
        self.window_icon = load_window_icon(self, 32)
        for ico_path in ICO_PATHS:
            if ico_path.exists():
                try:
                    self.iconbitmap(str(ico_path))
                except Exception as exc:
                    print(f"Не удалось установить iconbitmap {ico_path.name}: {exc}")
                break
        if self.window_icon:
            self.iconphoto(True, self.window_icon)
        self.cells, _skipped = load_db()
        if not self.cells:
            messagebox.showerror("База пуста", "cells_db.txt отсутствует или не содержит корректных ячеек.")
            self.after(100, self.destroy)
            return
        self.cell_by_label = {self.cell_label(c): c for c in self.cells}
        self.entries: dict[str, ctk.CTkEntry] = {}
        self.errors: dict[str, ctk.CTkLabel] = {}
        self.field_widgets: dict[str, tuple[ctk.CTkLabel, ctk.CTkEntry, ctk.CTkLabel]] = {}
        self.dim_extra_widgets: list[ctk.CTkBaseClass] = []
        self.packs: list[Pack] = []
        self.selected: Optional[Pack] = None
        self.calc_running = False
        self.calc_anim_job: Optional[str] = None
        self.calc_anim_step = 0
        self._build()
        self.restore_state()
        self.on_cell_change(self.cell_combo.get())
        self.protocol("WM_DELETE_WINDOW", self.on_close)
        signal.signal(signal.SIGINT, lambda _sig, _frame: self.after(0, self.on_close))
        self.after(100, self.keep_signal_handling_alive)

    def _center(self, wf: float, hf: float) -> None:
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        w, h = int(sw * wf), int(sh * hf)
        self.geometry(f"{w}x{h}+{(sw - w) // 2}+{(sh - h) // 2}")

    def cell_label(self, cell: Cell) -> str:
        return f"{cell.name} — {cell.brand} — {fmt(cell.capacity, 'А·ч')} — Uном {fmt(cell.voltage_nom, 'В')}"

    def _build(self) -> None:
        self.page = ctk.CTkScrollableFrame(self, fg_color=BG)
        self.page.pack(fill="both", expand=True)

        header = ctk.CTkFrame(self.page, fg_color=BG)
        header.pack(fill="x", padx=16, pady=(14, 6))
        if self.logo:
            ctk.CTkLabel(header, image=self.logo, text="").pack(side="left")
        ctk.CTkLabel(header, text="RD конфигуратор АКБ", font=("GOST type A", 26, "bold")).pack(side="left", padx=12)

        top = ctk.CTkFrame(self.page, fg_color=BG)
        top.pack(fill="x", padx=16, pady=8)
        top.grid_columnconfigure(0, weight=1)
        top.grid_columnconfigure(1, weight=1)

        inputs = ctk.CTkFrame(top, fg_color=BG)
        inputs.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
        self.cell_combo = ctk.CTkOptionMenu(inputs, values=list(self.cell_by_label), command=self.on_cell_change, fg_color=FIELD_BG, button_color=BUTTON, text_color="#111", font=FONT)
        self.cell_combo.pack(fill="x", pady=(0, 8))
        fields = [
            ("load", "Нагрузка, Вт", ""),
            ("time", "Требуемое время работы, ч", ""),
            ("depth", "Длина сборки, мм", ""),
            ("width", "Ширина сборки, мм", ""),
            ("height", "Высота сборки, мм", ""),
            ("cut", "Допуск габаритных размеров, %", "0"),
            ("wall", "Толщина стенки корпуса, мм", "0"),
            ("gap", "Зазор между ячейками, мм", "2"),
            ("bms", "Место под BMS, мм", "0"),
        ]
        self.consider_var = tk.BooleanVar(value=False)
        self.consider_box = ctk.CTkCheckBox(inputs, text="Учитывать габариты", variable=self.consider_var, command=self.on_consider_toggle, font=FONT)
        self.consider_box.pack(anchor="w", pady=(0, 6))
        self.dim_group = ctk.CTkFrame(inputs, fg_color="#fff4d8")
        ctk.CTkLabel(self.dim_group, text="Габариты, зазоры и допуски", anchor="w", font=("GOST type A", 15, "bold")).pack(fill="x", padx=10, pady=(8, 4))
        self.bms_group = ctk.CTkFrame(inputs, fg_color="#eaf2ff")
        ctk.CTkLabel(self.bms_group, text="Место под BMS", anchor="w", font=("GOST type A", 15, "bold")).pack(fill="x", padx=10, pady=(8, 4))
        for key, label, default in fields:
            if key in {"depth", "width", "height", "cut", "wall", "gap"}:
                parent = self.dim_group
                padx = 10
            elif key == "bms":
                parent = self.bms_group
                padx = 10
            else:
                parent = inputs
                padx = 0
            label_widget = ctk.CTkLabel(parent, text=label, anchor="w", font=FONT)
            label_widget.pack(fill="x", padx=padx)
            entry = ctk.CTkEntry(parent, width=230, fg_color=FIELD_BG, border_width=0, font=FONT)
            entry.insert(0, default)
            entry.pack(anchor="w", padx=padx)
            err = ctk.CTkLabel(parent, text="", text_color=ERROR, anchor="w", height=15, font=("GOST type A", 11))
            err.pack(fill="x", padx=padx)
            self.entries[key] = entry
            self.errors[key] = err
            self.field_widgets[key] = (label_widget, entry, err)
            if key == "bms":
                self.bms_side_label = ctk.CTkLabel(self.bms_group, text="Сторона BMS", anchor="w", font=FONT)
                self.bms_side_label.pack(fill="x", padx=10)
                self.bms_side_menu = ctk.CTkOptionMenu(
                    self.bms_group,
                    width=230,
                    values=["спереди", "сбоку", "сверху"],
                    fg_color=FIELD_BG,
                    button_color=BUTTON,
                    button_hover_color=BUTTON_HOVER,
                    text_color="#111111",
                    font=FONT,
                    command=lambda _value: self.save_state(),
                )
                self.bms_side_menu.set("спереди")
                self.bms_side_menu.pack(anchor="w", padx=10, pady=(0, 10))
                self.dim_extra_widgets.extend([self.bms_side_label, self.bms_side_menu])

        self.mirror = ctk.CTkFrame(top, fg_color="#ffffff", border_width=1, border_color="#d0d0d0")
        self.mirror.grid(row=0, column=1, sticky="nsew")

        self.calc_button = ctk.CTkButton(self.page, text="Рассчитать сборку", height=48, fg_color=BUTTON, hover_color=BUTTON_HOVER, text_color="#111", font=("GOST type A", 20, "bold"), command=self.run_calc)
        self.calc_button.pack(fill="x", padx=16, pady=8)
        self.calc_status_frame = ctk.CTkFrame(self.page, fg_color="#eef3f6", border_width=1, border_color="#c7d8df")
        self.calc_status_frame.grid_columnconfigure(0, weight=1)
        self.calc_status_label = ctk.CTkLabel(
            self.calc_status_frame,
            text="",
            anchor="w",
            text_color="#111111",
            fg_color="#eef3f6",
            font=("GOST type A", 18, "bold"),
        )
        self.calc_status_label.grid(row=0, column=0, sticky="ew", padx=12, pady=(8, 4))
        self.calc_progress = ctk.CTkProgressBar(
            self.calc_status_frame,
            mode="indeterminate",
            height=12,
            fg_color="#d5e2e7",
            progress_color=BUTTON,
        )
        self.calc_progress.grid(row=1, column=0, sticky="ew", padx=12, pady=(0, 10))

        self.output = ctk.CTkFrame(self.page, height=96, fg_color="#ffffff", border_width=1, border_color="#c8c8c8")
        self.output.pack(fill="x", padx=16, pady=(0, 16))
        self.output.pack_propagate(False)
        self.summary_row = ctk.CTkFrame(self.output, height=130, fg_color="#ffffff")
        self.summary_row.grid_columnconfigure(0, weight=2)
        self.summary_row.grid_columnconfigure(1, weight=1)
        self.summary_row.grid_rowconfigure(0, weight=1)
        self.variant_bar = ctk.CTkFrame(self.summary_row, fg_color="#ffffff")
        self.variant_bar.grid(row=0, column=0, sticky="nsew", padx=(8, 6), pady=8)
        self.description = ctk.CTkLabel(
            self.summary_row,
            text="",
            fg_color="#ffffff",
            text_color="#111111",
            anchor="nw",
            justify="left",
            font=("GOST type A", 20, "bold"),
            wraplength=360,
            height=104,
        )
        self.description.grid(row=0, column=1, sticky="nsew", padx=(6, 8), pady=8)
        self.runtime_label = ctk.CTkLabel(
            self.summary_row,
            text="",
            fg_color="#ffffff",
            text_color="#111111",
            anchor="w",
            justify="left",
            font=("GOST type A", 22, "bold"),
            height=30,
        )
        self.runtime_label.grid(row=1, column=1, sticky="ew", padx=(6, 8), pady=(0, 8))

        self.empty_output_label = ctk.CTkLabel(
            self.output,
            text="Вывод появится здесь после расчёта сборки.",
            fg_color="#ffffff",
            text_color="#777777",
            font=("GOST type A", 20),
            height=28,
        )

        self.detail = ctk.CTkFrame(self.output, height=630, fg_color="#ffffff")
        self.detail.pack_propagate(False)
        self.detail.grid_columnconfigure(0, weight=1)
        self.detail.grid_rowconfigure(0, weight=1)
        self.detail.grid_rowconfigure(1, weight=1)
        front_panel = ctk.CTkFrame(self.detail, fg_color="#ffffff")
        front_panel.grid(row=0, column=0, sticky="nsew", pady=(0, 6))
        front_panel.grid_rowconfigure(1, weight=1)
        front_panel.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(front_panel, text="Вид спереди", anchor="w", font=("GOST type A", 18, "bold")).grid(row=0, column=0, sticky="ew", pady=(0, 4))
        self.front_canvas = tk.Canvas(front_panel, bg="white", highlightthickness=1, highlightbackground="#e0e0e0")
        self.front_canvas.grid(row=1, column=0, sticky="nsew")

        side_panel = ctk.CTkFrame(self.detail, fg_color="#ffffff")
        side_panel.grid(row=1, column=0, sticky="nsew")
        side_panel.grid_rowconfigure(1, weight=1)
        side_panel.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(side_panel, text="Вид сбоку", anchor="w", font=("GOST type A", 18, "bold")).grid(row=0, column=0, sticky="ew", pady=(0, 4))
        self.side_canvas = tk.Canvas(side_panel, bg="white", highlightthickness=1, highlightbackground="#e0e0e0")
        self.side_canvas.grid(row=1, column=0, sticky="nsew")
        self.export_button = ctk.CTkButton(
            self.output,
            text="Экспорт в Excel",
            height=54,
            fg_color="#1b7a1b",
            hover_color="#146414",
            text_color="#ffffff",
            font=("GOST type A", 20, "bold"),
            command=self.export_excel,
        )
        self.toggle_cut()
        self.show_empty_output()
        self.bind_all("<Return>", self.on_calc_hotkey)
        self.bind_all("<F5>", self.on_calc_hotkey)

    def toggle_cut(self) -> None:
        dim_keys = {"depth", "width", "height", "cut", "wall", "gap"}
        if self.consider_var.get():
            if not self.dim_group.winfo_ismapped():
                self.dim_group.pack(fill="x", pady=(4, 8), padx=0, after=self.consider_box)
            if not self.bms_group.winfo_ismapped():
                self.bms_group.pack(fill="x", pady=(0, 8), padx=0, after=self.dim_group)
        else:
            for key in dim_keys:
                self.errors[key].configure(text="")
            self.dim_group.pack_forget()
            if not self.bms_group.winfo_ismapped():
                self.bms_group.pack(fill="x", pady=(4, 8), padx=0, after=self.consider_box)

    def on_consider_toggle(self) -> None:
        self.toggle_cut()
        self.save_state()

    def restore_state(self) -> None:
        if not STATE_PATH.exists():
            return
        try:
            state = json.loads(STATE_PATH.read_text(encoding="utf-8"))
        except Exception as exc:
            print(f"Не удалось прочитать настройки {STATE_PATH.name}: {exc}")
            return
        cell_label = state.get("cell")
        if isinstance(cell_label, str) and cell_label in self.cell_by_label:
            self.cell_combo.set(cell_label)
        self.consider_var.set(bool(state.get("consider_dims", False)))
        fields = state.get("fields", {})
        if isinstance(fields, dict):
            for key, value in fields.items():
                if key in self.entries:
                    self.entries[key].delete(0, "end")
                    self.entries[key].insert(0, str(value))
        bms_side = state.get("bms_side")
        if isinstance(bms_side, str) and bms_side in {"спереди", "сбоку", "сверху"}:
            self.bms_side_menu.set(bms_side)
        self.toggle_cut()

    def save_state(self) -> None:
        if not self.entries:
            return
        state = {
            "cell": self.cell_combo.get(),
            "consider_dims": bool(self.consider_var.get()),
            "bms_side": self.bms_side_menu.get(),
            "fields": {key: entry.get() for key, entry in self.entries.items()},
        }
        try:
            STATE_PATH.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
        except Exception as exc:
            print(f"Не удалось сохранить настройки {STATE_PATH.name}: {exc}")

    def on_close(self) -> None:
        self.save_state()
        self.destroy()

    def keep_signal_handling_alive(self) -> None:
        self.after(100, self.keep_signal_handling_alive)

    def on_calc_hotkey(self, _event: tk.Event) -> str:
        self.run_calc()
        return "break"

    def show_empty_output(self) -> None:
        self.output.configure(height=96)
        self.summary_row.pack_forget()
        self.detail.pack_forget()
        self.export_button.pack_forget()
        self.empty_output_label.pack(fill="both", expand=True, padx=8, pady=8)
        self.description.configure(text="")
        self.runtime_label.configure(text="")
        self.front_canvas.delete("all")
        self.side_canvas.delete("all")

    def show_detail_output(self) -> None:
        self.output.configure(height=980)
        self.empty_output_label.pack_forget()
        if not self.summary_row.winfo_ismapped():
            self.summary_row.pack(fill="x", padx=8, pady=(8, 4))
        if not self.detail.winfo_ismapped():
            self.detail.pack(fill="both", expand=True, padx=8, pady=8)
        if not self.export_button.winfo_ismapped():
            self.export_button.pack(fill="x", padx=8, pady=(0, 10))

    def scroll_to_output(self) -> None:
        canvas = getattr(self.page, "_parent_canvas", None)
        if canvas is not None:
            canvas.yview_moveto(0.55)

    def on_cell_change(self, _value: str) -> None:
        self.save_state()
        cell = self.current_cell()
        for widget in self.mirror.winfo_children():
            widget.destroy()
        rows = [
            ("Наименование", cell.name), ("Бренд", cell.brand), ("Тип химии", cell.chemistry),
            ("Тип формы", cell.shape), ("Габариты Д×Ш×В, мм", f"{cell.length:g}×{cell.width:g}×{cell.height:g}"),
            ("Ёмкость", fmt(cell.capacity, "А·ч")), ("Напряжение ном.", fmt(cell.voltage_nom, "В")),
            ("Напряжение мин./макс.", f"{fmt(cell.voltage_min, 'В')} / {fmt(cell.voltage_max, 'В')}"),
            ("Макс. ток разряда", fmt(cell.max_discharge, "А")), ("Макс. ток заряда", fmt(cell.max_charge, "А")),
            ("Мин./Макс. темп.", f"{fmt(cell.temp_min, '°C')} / {fmt(cell.temp_max, '°C')}"),
            ("Масса", fmt(cell.mass, "г")), ("Внутр. сопротивление", fmt(cell.resistance, "мОм")),
            ("Циклов", fmt(cell.cycles, "")), ("После 1000 циклов", fmt(cell.capacity_1000, "%")),
        ]
        for key, value in rows:
            line = ctk.CTkFrame(self.mirror, fg_color="#ffffff")
            line.pack(fill="x", padx=10, pady=2)
            ctk.CTkLabel(line, text=f"{key} →", width=180, anchor="w", font=FONT).pack(side="left")
            ctk.CTkLabel(line, text=value, anchor="w", font=FONT).pack(side="left", fill="x", expand=True)

    def current_cell(self) -> Cell:
        return self.cell_by_label[self.cell_combo.get()]

    def parse_input(self, key: str, required: bool = True) -> Optional[float]:
        self.errors[key].configure(text="")
        text = self.entries[key].get().strip().replace(",", ".")
        if not text and required:
            self.errors[key].configure(text="поле обязательно")
            return None
        if not text:
            return 0.0
        try:
            return float(text)
        except ValueError:
            self.errors[key].configure(text="должно быть число")
            return None

    def start_calc_animation(self) -> None:
        self.calc_running = True
        self.calc_anim_step = 0
        if not self.calc_status_frame.winfo_ismapped():
            self.calc_status_frame.pack(fill="x", padx=16, pady=(0, 8), after=self.calc_button)
        self.calc_status_label.configure(text="Идёт расчёт   [■□□□□]")
        self.calc_progress.start()
        self.calc_button.configure(text="Идёт расчёт...", command=lambda: None)
        self.animate_calc_button()

    def animate_calc_button(self) -> None:
        if not self.calc_running:
            return
        blocks = ("[■□□□□]", "[■■□□□]", "[■■■□□]", "[■■■■□]", "[■■■■■]")
        dots = "." * ((self.calc_anim_step % 3) + 1)
        self.calc_status_label.configure(text=f"Идёт расчёт{dots}   {blocks[self.calc_anim_step % len(blocks)]}")
        self.calc_button.configure(text="Идёт расчёт...")
        self.calc_anim_step += 1
        self.calc_anim_job = self.after(360, self.animate_calc_button)

    def stop_calc_animation(self) -> None:
        self.calc_running = False
        try:
            self.calc_progress.stop()
        except tk.TclError:
            pass
        if self.calc_status_frame.winfo_ismapped():
            self.calc_status_frame.pack_forget()
        if self.calc_anim_job is not None:
            try:
                self.after_cancel(self.calc_anim_job)
            except tk.TclError:
                pass
            self.calc_anim_job = None
        self.calc_button.configure(text="Рассчитать сборку", command=self.run_calc)

    def run_calc(self, keep_selection: bool = False) -> None:
        if self.calc_running:
            return
        self.save_state()
        values = {k: self.parse_input(k, k in {"load", "time"} or self.consider_var.get() and k in {"depth", "width", "height"}) for k in self.entries}
        if any(v is None for v in values.values()):
            return
        if values["load"] <= 0 or values["time"] <= 0:
            self.errors["load"].configure(text="нагрузка и время должны быть больше 0")
            return
        cell = self.current_cell()
        consider_dims = bool(self.consider_var.get())
        bms_side = {"спереди": "front", "сбоку": "side", "сверху": "top"}.get(self.bms_side_menu.get(), "front")
        previous = self.selected if keep_selection else None
        self.start_calc_animation()
        threading.Thread(target=self._calc_worker, args=(cell, values, consider_dims, bms_side, previous), daemon=True).start()

    def _calc_worker(
        self,
        cell: Cell,
        values: dict[str, float],
        consider_dims: bool,
        bms_side: str,
        previous: Optional[Pack],
    ) -> None:
        try:
            packs = calculate(
                cell, values["load"], values["time"], consider_dims,
                (values["depth"], values["width"], values["height"]),
                values["cut"],
                values["wall"],
                values["gap"],
                values["bms"],
                bms_side,
            )
            error = ""
        except Exception as exc:
            packs = []
            error = str(exc)
        try:
            self.after(0, lambda: self.finish_calc(packs, previous, error))
        except RuntimeError:
            pass

    def finish_calc(self, packs: list[Pack], previous: Optional[Pack], error: str) -> None:
        self.stop_calc_animation()
        if error:
            self.description.configure(text=f"Ошибка расчёта: {error}")
            self.show_empty_output()
            return
        self.packs = packs
        if self.consider_var.get() and self.packs and not any(p.ok_dims for p in self.packs):
            self.description.configure(text="Ни одна сборка не влезает в заданные габариты.\nПоказаны ближайшие варианты.")
        self.draw_variant_buttons()
        if self.packs:
            selected = previous if previous in self.packs else self.packs[0]
            self.select_pack(selected)
            self.after(120, self.scroll_to_output)
        else:
            self.show_empty_output()

    def draw_variant_buttons(self) -> None:
        for widget in self.variant_bar.winfo_children():
            widget.destroy()
        colors = self.variant_colors()
        for col in range(5):
            self.variant_bar.grid_columnconfigure(col, weight=1)
        for index, (pack, color) in enumerate(zip(self.packs, colors)):
            layout = f"{pack.grid[0]}×{pack.grid[1]}×{pack.grid[2]}"
            label = f"{pack.s}S{pack.p}P\n{layout}"
            hover = self.darken_color(color)
            ctk.CTkButton(
                self.variant_bar,
                text=label,
                width=95,
                height=52,
                fg_color=color,
                hover_color=hover,
                text_color="#111",
                command=lambda p=pack: self.select_pack(p),
            ).grid(row=index // 5, column=index % 5, padx=4, pady=4, sticky="ew")

    def variant_colors(self) -> list[str]:
        return [self.pack_color(pack) for pack in self.packs]

    def reserve_color(self, reserve_ratio: float) -> str:
        # 0% запаса - желтый, 50% и больше - уверенно зеленый.
        ratio = max(0.0, min(1.0, reserve_ratio / 0.5))
        start = (242, 201, 76)
        end = (49, 158, 80)
        rgb = tuple(round(start[i] + (end[i] - start[i]) * ratio) for i in range(3))
        return "#{:02x}{:02x}{:02x}".format(*rgb)

    def pack_color(self, pack: Pack) -> str:
        if not pack.ok_time or not pack.ok_dims:
            return "#d9534f"
        required = self.parse_input("time", False) or 0.0
        if required <= 0:
            return "#f2c94c"
        reserve_ratio = (pack.runtime - required) / required
        return self.reserve_color(reserve_ratio)

    def darken_color(self, color: str) -> str:
        color = color.lstrip("#")
        rgb = [int(color[i:i + 2], 16) for i in (0, 2, 4)]
        dark = [max(0, int(v * 0.82)) for v in rgb]
        return "#{:02x}{:02x}{:02x}".format(*dark)

    def select_pack(self, pack: Pack) -> None:
        self.show_detail_output()
        self.selected = pack
        cell = self.current_cell()
        selected_color = self.pack_color(pack)
        orient = "стоит" if pack.orientation[2] == "L" else "лежит"
        text = [
            f"{cell.name}",
            f"Тип химии: {cell.chemistry}",
            f"Форма ячейки: {shape_name(cell.shape)}, ориентация {orient} ({''.join(pack.orientation)})",
            f"{pack.s}S × {pack.p}P",
            f"ЧИСЛО ИСПОЛЬЗУЕМЫХ ЯЧЕЕК = {pack.n}",
            f"Общая ёмкость: {fmt(pack.capacity, 'А·ч')}",
            f"Номинальное напряжение сборки: {fmt(pack.voltage, 'В')}",
            f"Общая энергия: {fmt(pack.energy, 'Вт·ч')}",
            f"Макс ток разряда сборки: {fmt(pack.max_current, 'А')}",
            f"Температура: {fmt(cell.temp_min, '°C')} / {fmt(cell.temp_max, '°C')}",
            f"Ёмкость после 1000 циклов: {fmt(cell.capacity_1000, '%')}",
            f"Габариты Д×Ш×В: {pack.dims[0]:.1f}×{pack.dims[1]:.1f}×{pack.dims[2]:.1f} мм",
        ]
        if pack.warning:
            text.append(pack.warning)
        self.description.configure(text="\n".join(text))
        self.runtime_label.configure(text=f"Расчётное время: {fmt(pack.runtime, 'ч')}", text_color=selected_color)
        self.update_idletasks()
        self.draw_pack(pack)
        self.after(80, lambda p=pack: self.draw_pack(p) if self.selected is p else None)

    def draw_pack(self, pack: Pack) -> None:
        self._draw_canvas(self.front_canvas, pack, front=True)
        self._draw_canvas(self.side_canvas, pack, front=False)

    def _draw_canvas(self, canvas: tk.Canvas, pack: Pack, front: bool) -> None:
        canvas.delete("all")
        canvas.update_idletasks()
        w, h = max(240, canvas.winfo_width()), max(180, canvas.winfo_height())
        d, wid, hei = pack.dims
        view_w, view_h = (wid, hei) if front else (d, hei)
        target = (0.0, 0.0, 0.0)
        cut = 0.0
        target_view_w, target_view_h = 0.0, 0.0
        cut_view_w, cut_view_h = 0.0, 0.0
        if self.consider_var.get():
            target = tuple(self.parse_input(k, False) or 0.0 for k in ("depth", "width", "height"))
            cut = self.parse_input("cut", False) or 0.0
            target_view_w, target_view_h = (target[1], target[2]) if front else (target[0], target[2])
            cut_target = tuple(v * (1 + cut / 100.0) for v in target)
            cut_view_w, cut_view_h = (cut_target[1], cut_target[2]) if front else (cut_target[0], cut_target[2])
        scale_w = max(view_w, target_view_w, cut_view_w, 1)
        scale_h = max(view_h, target_view_h, cut_view_h, 1)
        left_pad = max(24, min(48, int(w * 0.05)))
        top_pad = max(20, min(38, int(h * 0.07)))
        right_pad = max(54, min(82, int(w * 0.10)))
        bottom_pad = max(44, min(70, int(h * 0.14)))
        scale = min((w - left_pad - right_pad) / scale_w, (h - top_pad - bottom_pad) / scale_h)
        ox, oy = left_pad, top_pad
        rw, rh = view_w * scale, view_h * scale
        canvas.create_rectangle(ox, oy, ox + rw, oy + rh, outline="#cc0000", width=3)
        bms = self.parse_input("bms", False) or 0.0
        bms_side = self.bms_side_menu.get()
        cell_view_w, cell_view_h = view_w, view_h
        bms_rect = None
        if self.consider_var.get() and bms > 0:
            bms_area_w = target_view_w if target_view_w > 0 else view_w
            bms_area_h = target_view_h if target_view_h > 0 else view_h
            bms_area_px = bms_area_w * scale
            bms_area_py = bms_area_h * scale
            bms_px = min(bms_area_px, bms * scale)
            bms_py = min(bms_area_py, bms * scale)
            if bms_side == "спереди" and not front:
                cell_view_w = max(0.0, view_w - bms)
                bms_rect = (ox + bms_area_px - bms_px, oy, ox + bms_area_px, oy + bms_area_py)
            elif bms_side == "сбоку" and front:
                cell_view_w = max(0.0, view_w - bms)
                bms_rect = (ox + bms_area_px - bms_px, oy, ox + bms_area_px, oy + bms_area_py)
            elif bms_side == "сверху":
                cell_view_h = max(0.0, view_h - bms)
                bms_rect = (ox, oy, ox + bms_area_px, oy + bms_py)
        cell_rw, cell_rh = cell_view_w * scale, cell_view_h * scale
        gx, gy, gz = pack.grid
        cols, rows = (gy, gz) if front else (gx, gz)
        cell_origin_y = oy + (rh - cell_rh if bms_rect and bms_side == "сверху" else 0)
        cell_w, cell_h = cell_rw / max(cols, 1), cell_rh / max(rows, 1)
        for y in range(rows):
            for x in range(cols):
                x0 = ox + x * cell_w + 2
                y0 = cell_origin_y + y * cell_h + 2
                x1 = ox + (x + 1) * cell_w - 2
                y1 = cell_origin_y + (y + 1) * cell_h - 2
                if self.current_cell().shape == "C" and front and pack.orientation[2] == "L":
                    canvas.create_oval(x0, y0, x1, y1, outline="#777777", width=3)
                else:
                    canvas.create_rectangle(x0, y0, x1, y1, outline="#777777", width=3)
        if bms_rect:
            self.draw_hatched_rect(canvas, bms_rect)
            self.draw_bms_label(canvas, bms_rect)
        if self.consider_var.get():
            if cut > 0:
                canvas.create_rectangle(
                    ox,
                    oy,
                    ox + cut_view_w * scale,
                    oy + cut_view_h * scale,
                    outline="#d6a400",
                    width=4,
                    dash=(10, 5),
                )
            canvas.create_rectangle(
                ox,
                oy,
                ox + target_view_w * scale,
                oy + target_view_h * scale,
                outline="#159947",
                width=6,
            )
            if target_view_w > 0 and target_view_h > 0:
                self.draw_dimension_labels(
                    canvas,
                    ox,
                    oy,
                    target_view_w * scale,
                    target_view_h * scale,
                    target_view_w,
                    target_view_h,
                    "#159947",
                    outside=False,
                )
        self.draw_dimension_labels(canvas, ox, oy, rw, rh, view_w, view_h, "#ffffff", outside=True)

    def dimension_font_size(self, side_px: float, text: str, limit: int) -> int:
        by_side = int(side_px / max(len(text) * 0.62, 1))
        return max(9, min(limit, by_side))

    def draw_dimension_labels(
        self,
        canvas: tk.Canvas,
        x: float,
        y: float,
        rect_w: float,
        rect_h: float,
        value_w: float,
        value_h: float,
        color: str,
        outside: bool,
    ) -> None:
        text_w = f"{value_w:.1f} мм"
        text_h = f"{value_h:.1f} мм"
        font_w = self.dimension_font_size(rect_w, text_w, 36 if outside else 28)
        font_h = self.dimension_font_size(rect_h, text_h, 36 if outside else 28)
        bg = "#6a6a6a"
        if outside:
            gap = 12
            self.draw_label_box(canvas, x + rect_w / 2, y + rect_h + gap + font_w * 0.45, text_w, font_w, color, bg, angle=0)
            self.draw_label_box(canvas, x + rect_w + gap + font_h * 0.45, y + rect_h / 2, text_h, font_h, color, bg, angle=90)
        else:
            gap = 12
            self.draw_label_box(canvas, x + rect_w / 2, y + gap + font_w * 0.45, text_w, font_w, color, bg, angle=0)
            self.draw_label_box(canvas, x + rect_w - gap - font_h * 0.45, y + rect_h / 2, text_h, font_h, color, bg, angle=90)

    def draw_label_box(
        self,
        canvas: tk.Canvas,
        cx: float,
        cy: float,
        text: str,
        font_size: int,
        color: str,
        bg: str,
        angle: int,
    ) -> None:
        color = self.contrast_text_color(bg)
        pad_x = max(10, int(font_size * 0.45))
        pad_y = max(6, int(font_size * 0.28))
        text_w = max(28, len(text) * font_size * 0.55)
        text_h = font_size * 1.1
        text_w *= 1.15
        if angle == 90:
            box_w, box_h = text_h + pad_y * 2, text_w + pad_x * 2
        else:
            box_w, box_h = text_w + pad_x * 2, text_h + pad_y * 2
        canvas.create_rectangle(cx - box_w / 2, cy - box_h / 2, cx + box_w / 2, cy + box_h / 2, fill=bg, outline=bg)
        canvas.create_text(cx, cy, text=text, fill=color, font=("GOST type A", font_size), angle=angle, anchor="center")

    def contrast_text_color(self, bg: str) -> str:
        bg = bg.lstrip("#")
        r, g, b = (int(bg[i:i + 2], 16) for i in (0, 2, 4))
        luminance = 0.2126 * r + 0.7152 * g + 0.0722 * b
        return "#ffffff" if luminance < 150 else "#111111"

    def draw_hatched_rect(self, canvas: tk.Canvas, rect: tuple[float, float, float, float]) -> None:
        x0, y0, x1, y1 = rect
        canvas.create_rectangle(x0, y0, x1, y1, outline="#555555", fill="#e9f5ea", width=2)
        step = 12
        height = y1 - y0
        pos = x0 - height
        while pos < x1:
            raw_a = (pos, y1)
            raw_b = (pos + height, y0)
            clipped = self.clip_line_to_rect(raw_a, raw_b, rect)
            if clipped:
                (xa, ya), (xb, yb) = clipped
                canvas.create_line(xa, ya, xb, yb, fill="#cc0000", width=2)
            pos += step
        canvas.create_rectangle(x0, y0, x1, y1, outline="#555555", width=2)

    def draw_bms_label(self, canvas: tk.Canvas, rect: tuple[float, float, float, float]) -> None:
        x0, y0, x1, y1 = rect
        width, height = x1 - x0, y1 - y0
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        angle = 90 if height > width * 1.4 else 0
        text = "BMS"
        pad_x, pad_y = 8, 5
        if angle == 90:
            max_font_by_width = int((width - pad_x * 2) / 1.25)
            max_font_by_height = int((height - pad_y * 2) / (len(text) * 0.72))
        else:
            max_font_by_width = int((width - pad_x * 2) / (len(text) * 0.72))
            max_font_by_height = int((height - pad_y * 2) / 1.25)
        font_size = max(8, min(28, max_font_by_width, max_font_by_height))
        text_w = max(20, len(text) * font_size * 0.72)
        text_h = font_size * 1.25
        if angle == 90:
            box_w = min(width - 4, text_h + pad_y * 2)
            box_h = min(height - 4, text_w + pad_x * 2)
        else:
            box_w = min(width - 4, text_w + pad_x * 2)
            box_h = min(height - 4, text_h + pad_y * 2)
        canvas.create_rectangle(
            cx - box_w / 2,
            cy - box_h / 2,
            cx + box_w / 2,
            cy + box_h / 2,
            fill="#ffffff",
            outline="#ffffff",
        )
        canvas.create_text(cx, cy, text=text, angle=angle, fill="#111111", font=("GOST type A", font_size, "bold"), anchor="center")

    def clip_line_to_rect(
        self,
        p0: tuple[float, float],
        p1: tuple[float, float],
        rect: tuple[float, float, float, float],
    ) -> Optional[tuple[tuple[float, float], tuple[float, float]]]:
        x0, y0 = p0
        x1, y1 = p1
        xmin, ymin, xmax, ymax = rect
        dx, dy = x1 - x0, y1 - y0
        t0, t1 = 0.0, 1.0
        for p, q in ((-dx, x0 - xmin), (dx, xmax - x0), (-dy, y0 - ymin), (dy, ymax - y0)):
            if p == 0:
                if q < 0:
                    return None
                continue
            r = q / p
            if p < 0:
                if r > t1:
                    return None
                t0 = max(t0, r)
            else:
                if r < t0:
                    return None
                t1 = min(t1, r)
        return ((x0 + t0 * dx, y0 + t0 * dy), (x0 + t1 * dx, y0 + t1 * dy))

    def export_excel(self) -> None:
        if not self.packs:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            messagebox.showerror("Нет openpyxl", "Установите зависимость: pip install openpyxl")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], initialfile="rd_pack_variants.xlsx")
        if not path:
            return
        wb = Workbook()
        wb.remove(wb.active)
        cell = self.current_cell()
        inputs = {key: self.entries[key].get() for key in self.entries}

        title_fill = PatternFill("solid", fgColor="1B7A1B")
        section_fill = PatternFill("solid", fgColor="D9EAD3")
        warning_fill = PatternFill("solid", fgColor="F4CCCC")
        title_font = Font(bold=True, color="FFFFFF", size=14)
        section_font = Font(bold=True)

        for index, pack in enumerate(self.packs, start=1):
            sheet_name = f"{pack.s}S{pack.p}P"
            if sheet_name in wb.sheetnames:
                sheet_name = f"{sheet_name}_{index}"
            ws = wb.create_sheet(sheet_name[:31])
            ws.column_dimensions["A"].width = 34
            ws.column_dimensions["B"].width = 42
            ws.column_dimensions["C"].width = 18

            def section(title: str) -> None:
                row = ws.max_row + 1
                ws.cell(row, 1, title)
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
                ws.cell(row, 1).font = section_font
                ws.cell(row, 1).fill = section_fill

            def item(key: str, value: object, unit: str = "") -> None:
                row = ws.max_row + 1
                ws.cell(row, 1, key)
                ws.cell(row, 2, value)
                ws.cell(row, 3, unit)

            ws.append([f"RD конфигуратор АКБ - сборка {pack.s}S{pack.p}P"])
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
            ws["A1"].font = title_font
            ws["A1"].fill = title_fill
            ws["A1"].alignment = Alignment(horizontal="center")

            section("Исходная ячейка")
            item("Наименование", cell.name)
            item("Бренд", cell.brand)
            item("Тип химии", cell.chemistry)
            item("Тип формы", cell.shape)
            item("Габариты ячейки Д×Ш×В", f"{cell.length:g}×{cell.width:g}×{cell.height:g}", "мм")
            item("Ёмкость ячейки", round(cell.capacity, 3), "А·ч")
            item("Номинальное напряжение ячейки", round(cell.voltage_nom, 3), "В")
            item("Мин./макс. напряжение ячейки", f"{fmt(cell.voltage_min, 'В')} / {fmt(cell.voltage_max, 'В')}")
            item("Макс. ток разряда ячейки", round(cell.max_discharge, 3), "А")
            item("Макс. ток заряда ячейки", "" if cell.max_charge is None else round(cell.max_charge, 3), "А")
            item("Рабочая температура", f"{fmt(cell.temp_min, '°C')} / {fmt(cell.temp_max, '°C')}")
            item("Масса", "" if cell.mass is None else round(cell.mass, 3), "г")
            item("Внутр. сопротивление", "" if cell.resistance is None else round(cell.resistance, 3), "мОм")
            item("Циклов", "" if cell.cycles is None else round(cell.cycles, 0))
            item("Ёмкость после 1000 циклов", "" if cell.capacity_1000 is None else round(cell.capacity_1000, 2), "%")

            section("Конфигурация сборки")
            item("Схема", f"{pack.s}S × {pack.p}P")
            item("Число используемых ячеек", pack.n, "шт")
            item("3D-укладка a×b×c", f"{pack.grid[0]}×{pack.grid[1]}×{pack.grid[2]}")
            item("Ориентация осей", " → ".join(pack.orientation))
            item("Характерная метка", pack.marks or "обычный вариант")

            section("Электрические параметры")
            item("Общая ёмкость", round(pack.capacity, 3), "А·ч")
            item("Номинальное напряжение", round(pack.voltage, 3), "В")
            item("Общая энергия", round(pack.energy, 3), "Вт·ч")
            item("Макс. ток разряда сборки", round(pack.max_current, 3), "А")
            item("Расчётное время работы", round(pack.runtime, 3), "ч")
            item("Заданная нагрузка", inputs.get("load", ""), "Вт")
            item("Требуемое время", inputs.get("time", ""), "ч")

            section("Габариты и ограничения")
            item("Габариты сборки Д×Ш×В", f"{pack.dims[0]:.1f}×{pack.dims[1]:.1f}×{pack.dims[2]:.1f}", "мм")
            item("Объём сборки", round(pack.volume, 1), "мм³")
            item("Учитывать габариты", "да" if self.consider_var.get() else "нет")
            item("Целевые габариты Д×Ш×В", f"{inputs.get('depth', '')}×{inputs.get('width', '')}×{inputs.get('height', '')}", "мм")
            item("Допуск габаритных размеров", inputs.get("cut", ""), "%")
            item("Толщина стенки корпуса", inputs.get("wall", ""), "мм")
            item("Зазор между ячейками", inputs.get("gap", ""), "мм")
            item("Место под BMS", inputs.get("bms", ""), "мм")
            item("Сторона установки BMS", self.bms_side_menu.get())

            section("Статус")
            status = "проходит" if pack.ok_time and pack.ok_dims else pack.warning
            item("Итог", status)
            if pack.warning:
                ws.cell(ws.max_row, 2).fill = warning_fill

            for row in ws.iter_rows():
                for cell_obj in row:
                    cell_obj.alignment = Alignment(vertical="top", wrap_text=True)
        wb.save(path)


if __name__ == "__main__":
    try:
        MainApp().mainloop()
    except KeyboardInterrupt:
        sys.exit(0)
